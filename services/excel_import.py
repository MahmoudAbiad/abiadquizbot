"""Bounded, deterministic XLSX quiz intake; never evaluates Excel formulas."""
import json
import math
import re
import zipfile

from openpyxl import load_workbook

MAX_QUESTIONS = 120
MAX_ROWS = 5000
MAX_COLUMNS = 40
MAX_ARCHIVE_BYTES = 30 * 1024 * 1024


class ExcelImportError(ValueError):
    """An actionable upload error that is safe to show in Telegram."""


def _text(value):
    return str(value).strip() if value is not None else ""


def _header(value):
    return re.sub(r"\s+", " ", _text(value).lower()).replace("أ", "ا").replace("إ", "ا")


def _grid(value):
    return (isinstance(value, list) and 0 < len(value) <= 50
            and all(isinstance(row, list) and 0 < len(row) <= 20 for row in value)
            and all(len(row) == len(value[0]) for row in value)
            and all(isinstance(cell, (str, int, float)) and not isinstance(cell, bool)
                    and len(str(cell)) <= 2000 for row in value for cell in row))


def _parse_question(cells, columns, options, row_number):
    def fail(reason):
        raise ExcelImportError(f"الصف {row_number}: {reason}")

    if any(cell.data_type in {"f", "e"} for cell in cells):
        fail("استبدل صيغ Excel أو أخطاء الخلايا بقيم نصية ثابتة.")
    def value(key):
        index = columns.get(key)
        return _text(cells[index].value) if index is not None else ""

    question = value("question")
    if not question or len(question) > 10000:
        fail("نص السؤال مفقود أو أطول من 10000 حرف.")
    raw_options = [_text(cells[index].value) for _, index in options]
    while raw_options and not raw_options[-1]:
        raw_options.pop()
    if not 2 <= len(raw_options) <= 10 or any(not opt or len(opt) > 4000 for opt in raw_options):
        fail("يلزم خياران إلى 10 خيارات متتالية وغير فارغة، كل خيار حتى 4000 حرف.")

    number, answer = value("correct_number"), value("answer")
    correct = None
    if number:
        try:
            numeric = float(number)
            if not math.isfinite(numeric) or not numeric.is_integer():
                raise ValueError
            correct = int(numeric) - 1
        except ValueError:
            fail("رقم الإجابة الصحيحة يجب أن يكون عدداً صحيحاً يبدأ من 1.")
        if not 0 <= correct < len(raw_options):
            fail("رقم الإجابة الصحيحة خارج نطاق الخيارات (الترقيم يبدأ من 1).")
    if answer:
        if correct is not None:
            if raw_options[correct] != answer:
                fail("نص الإجابة الصحيحة لا يطابق الخيار المشار إليه برقم الإجابة؛ صحّح أحدهما.")
        else:
            matches = [i for i, option in enumerate(raw_options) if option == answer]
            if len(matches) != 1:
                fail("نص الإجابة الصحيحة يجب أن يطابق خياراً واحداً، أو استخدم رقم الإجابة الصحيحة.")
            correct = matches[0]
    if correct is None:
        fail("الإجابة الصحيحة مفقودة؛ استخدم نسخة الأستاذ أو أضف عمود «رقم الإجابة الصحيحة».")
    result = {"question": question, "options": raw_options, "correct_option_id": correct}
    for key in ("explanation", "type"):
        if value(key):
            if len(value(key)) > 10000:
                fail("الشرح أو نوع السؤال أطول من الحد المسموح.")
            result[key] = value(key)
    if value("points"):
        try:
            points = float(value("points"))
            if not math.isfinite(points) or not 0 < points <= 1000:
                raise ValueError
        except ValueError:
            fail("الدرجة يجب أن تكون رقماً موجباً حتى 1000.")
        result["points"] = points
    for key in ("table", "matrices"):
        raw = value(key)
        if not raw:
            continue
        try:
            if len(raw) > 30000:
                raise ValueError
            structured = json.loads(raw)
            if key == "table":
                valid = (isinstance(structured, dict) and _grid(structured.get("rows"))
                         and isinstance(structured.get("headers", []), list)
                         and len(structured.get("headers", [])) <= 20
                         and all(isinstance(h, str) and len(h) <= 2000 for h in structured.get("headers", [])))
            else:
                valid = (isinstance(structured, list) and 0 < len(structured) <= 10
                         and all(isinstance(m, dict) and _grid(m.get("rows"))
                                 and isinstance(m.get("label", ""), str) and len(m.get("label", "")) <= 2000
                                 and m.get("bracket", "square") in {"square", "round", "bar"}
                                 for m in structured))
            if not valid:
                raise ValueError
        except (ValueError, TypeError, RecursionError):
            fail("بيانات الجدول أو المصفوفات غير صالحة؛ احتفظ بتنسيق JSON الموجود في تصدير البوت.")
        result[key] = structured
    result["is_math"] = bool(result.get("table") or result.get("matrices") or
                             re.search(r"\$[^$]+\$|\\(?:frac|sqrt|begin|sum|int)\b", " ".join([question, *raw_options])))
    return result


def read_quiz_xlsx(path, default_title="كويز من Excel"):
    """Return (title, questions); the preview sheet and embedded images are ignored."""
    workbook = None
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if (len(members) > 2000 or sum(item.file_size for item in members) > MAX_ARCHIVE_BYTES
                    or any(item.flag_bits & 1 for item in members)):
                raise ExcelImportError("ملف Excel كبير أو مضغوط أكثر من المسموح؛ قسّمه إلى ملفات أصغر.")
            if "xl/workbook.xml" not in archive.namelist():
                raise ExcelImportError("الملف ليس مصنف Excel صالحاً بصيغة .xlsx.")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if len(workbook.worksheets) > 20:
            raise ExcelImportError("الحد الأقصى 20 ورقة في ملف Excel.")
        aliases = {
            "question": ["نص السؤال", "السؤال", "Question", "Question Text"],
            "correct_number": ["رقم الإجابة الصحيحة", "Correct Option Number"],
            "answer": ["الإجابة الصحيحة", "Correct Answer"],
            "explanation": ["الشرح", "Explanation"], "type": ["نوع السؤال", "Question Type"],
            "points": ["الدرجة", "Points"], "table": ["الجدول (JSON)", "Table (JSON)"],
            "matrices": ["المصفوفات (JSON)", "Matrices (JSON)"],
        }
        lookup = {_header(alias): key for key, values in aliases.items() for alias in values}
        candidates = []
        for sheet in workbook:
            headers = next(sheet.iter_rows(min_row=1, max_row=1, max_col=MAX_COLUMNS), ())
            columns, options = {}, []
            for i, cell in enumerate(headers):
                name = _header(cell.value)
                key = lookup.get(name)
                match = re.fullmatch(r"(?:الخيار|خيار|option)\s*([0-9٠-٩]+)", name)
                if key:
                    if key in columns:
                        raise ExcelImportError(f"ورقة «{sheet.title}»: يوجد عمود مكرر: {cell.value}.")
                    columns[key] = i
                if match:
                    options.append((int(match[1]), i))
            if "question" in columns and options:
                candidates.append((sheet, columns, sorted(options)))
        if len(candidates) != 1:
            raise ExcelImportError("يلزم ورقة أسئلة واحدة بأعمدة «نص السؤال»، «الخيار 1»، «الخيار 2»، و«رقم الإجابة الصحيحة» (يبدأ من 1). يمكنك تعديل نسخة الأستاذ المصدّرة من البوت وإرسالها.")
        sheet, columns, options = candidates[0]
        if [n for n, _ in options] != list(range(1, len(options) + 1)):
            raise ExcelImportError("أعمدة الخيارات يجب أن تبدأ من «الخيار 1» بترقيم متتالٍ دون تكرار.")
        if not {"correct_number", "answer"} & columns.keys():
            raise ExcelImportError("الملف بلا إجابات صحيحة؛ أرسل نسخة الأستاذ أو أضف عمود «رقم الإجابة الصحيحة» بترقيم يبدأ من 1.")
        if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
            raise ExcelImportError("ورقة الأسئلة كبيرة جداً؛ احذف الصفوف والأعمدة الزائدة.")
        questions = []
        # Ignore untrusted worksheet dimensions: cap actual streamed XML rows too.
        sheet.reset_dimensions()
        for number, row in enumerate(sheet.iter_rows(min_row=2, max_row=MAX_ROWS + 1, max_col=MAX_COLUMNS), 2):
            if not any(cell.value is not None and _text(cell.value) for cell in row):
                continue
            if number > MAX_ROWS or len(questions) >= MAX_QUESTIONS:
                raise ExcelImportError(f"الحد الأقصى {MAX_QUESTIONS} سؤالاً؛ قسّم الملف إلى اختبارات أصغر.")
            questions.append(_parse_question(row, columns, options, number))
        if not questions:
            raise ExcelImportError("ورقة الأسئلة فارغة؛ أضف أسئلة تحت صف العناوين.")
        title = default_title
        for name in ("معلومات الاختبار", "Quiz Info"):
            if name in workbook.sheetnames:
                info = workbook[name]
                label, cell = next(info.iter_rows(min_row=2, max_row=2, max_col=2))
                if _header(label.value) in {_header("عنوان الاختبار"), "quiz title"} and cell.data_type == "s" and _text(cell.value):
                    title = _text(cell.value)[:200]
        return title, questions
    except ExcelImportError:
        raise
    except Exception as exc:
        raise ExcelImportError("تعذّر قراءة ملف Excel. احفظه بصيغة .xlsx دون كلمة مرور ثم أعد إرساله.") from exc
    finally:
        if workbook is not None:
            workbook.close()
