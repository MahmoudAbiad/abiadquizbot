import asyncio
import io
import os
from pathlib import Path
import tempfile
import types
import unittest
from unittest.mock import AsyncMock, Mock, patch
import zipfile

from reliability_support import load_functions, FakeRedis, HTTPError
from services.redis_reliability import UserLock, complete_once, RELEASE
from services.upload_validation import validate_document, validate_image


class ClaimsTests(unittest.IsolatedAsyncioTestCase):
    async def test_concurrent_complete_starts_once(self):
        redis = FakeRedis()
        start = AsyncMock()
        await asyncio.gather(*(complete_once(redis, 'file', '1/source', start) for _ in range(20)))
        self.assertEqual(start.await_count, 1)

    async def test_failed_start_allows_retry(self):
        redis = FakeRedis()
        with self.assertRaises(ValueError):
            await complete_once(redis, 'audio', '1/source', AsyncMock(side_effect=ValueError()))
        start = AsyncMock()
        await complete_once(redis, 'audio', '1/source', start)
        start.assert_awaited_once()

    async def test_lock_two_instances_and_owned_release(self):
        redis = FakeRedis()
        first, second = UserLock(redis, 1), UserLock(redis, 1)
        self.assertTrue(await first.acquire())
        self.assertFalse(await second.acquire())
        await second.release()
        self.assertFalse(await second.acquire())
        await first.release()
        self.assertTrue(await second.acquire())
        await second.release()

    async def test_expired_owner_cannot_release_new_lock(self):
        redis = FakeRedis()
        first, second = UserLock(redis, 1, ttl=10), UserLock(redis, 1, ttl=10)
        self.assertTrue(await first.acquire())
        redis.time = 11
        self.assertTrue(await second.acquire())
        await first.release()
        self.assertEqual(redis.values[second.key], second.token)
        await second.release()

    async def test_redis_failure_does_not_start(self):
        redis = FakeRedis(); redis.fail = True
        start = AsyncMock()
        with self.assertRaises(ConnectionError):
            await complete_once(redis, 'file', '1/x', start)
        start.assert_not_awaited()


class UploadEndpointTests(unittest.IsolatedAsyncioTestCase):
    def env(self):
        from services.upload_validation import supported_extension, IMAGE_EXTENSIONS, DOCUMENT_EXTENSIONS, AUDIO_EXTENSIONS
        import re
        self.redis = FakeRedis()
        self.spawned = []
        def spawn(coro):
            self.spawned.append(coro)
            coro.close()
        async def once(kind, identity, start):
            return await complete_once(self.redis, kind, identity, start)
        return load_functions('webhook_server.py', {'_upload_extension', '_object_extension', 'image_upload_init', 'image_upload_complete', 'file_upload_init', 'file_upload_complete', 'audio_upload_complete'},
            HTTPException=HTTPError, os=os, re=re, supported_extension=supported_extension,
            IMAGE_EXTENSIONS=IMAGE_EXTENSIONS, DOCUMENT_EXTENSIONS=DOCUMENT_EXTENSIONS, AUDIO_EXTENSIONS=AUDIO_EXTENSIONS,
            bot=types.SimpleNamespace(token='fake'), verify_telegram_init_data=lambda *a, **k: (True, {'id':1}),
            FILE_UPLOAD_INIT_DATA_MAX_AGE_SECONDS=10, AUDIO_UPLOAD_INIT_DATA_MAX_AGE_SECONDS=10,
            FILE_UPLOAD_RATE_LIMIT_MAX_REQUESTS=20, FILE_UPLOAD_RATE_LIMIT_WINDOW_SECONDS=60,
            MAX_FILE_WEB_UPLOAD_SIZE=100, MAX_AUDIO_WEB_UPLOAD_SIZE=250,
            MAX_IMAGE_WEB_UPLOAD_COUNT=50, MAX_IMAGE_WEB_UPLOAD_SIZE_PER_IMAGE=15, MAX_IMAGE_WEB_UPLOAD_TOTAL_SIZE=500,
            _enforce_upload_rate_limit=AsyncMock(), _enforce_audio_upload_rate_limit=AsyncMock(),
            create_image_upload_targets=AsyncMock(return_value=[{'path':'x'}]), create_file_upload_target=AsyncMock(return_value={'path':'x'}),
            get_file_temp_object_size=AsyncMock(return_value=10), get_audio_temp_object_size=AsyncMock(return_value=10),
            delete_file_temp_batch=AsyncMock(), delete_file_temp=AsyncMock(), delete_audio_temp=AsyncMock(),
            process_web_uploaded_images=AsyncMock(), process_web_uploaded_file=AsyncMock(), process_web_uploaded_audio=AsyncMock(),
            _complete_once=once, spawn=spawn)

    async def test_init_total_and_per_file(self):
        env=self.env()
        for sizes in [[16], [15]*34]:
            with self.assertRaises(HTTPError) as exc:
                await env['image_upload_init'](types.SimpleNamespace(init_data='', file_sizes=sizes, file_names=['x.jpg']*len(sizes)))
            self.assertEqual(exc.exception.status_code,413)
        await env['image_upload_init'](types.SimpleNamespace(init_data='', file_sizes=[1]*50, file_names=['x.jpg']*50))
        await env['image_upload_init'](types.SimpleNamespace(init_data='', file_sizes=[10]*50, file_names=['x.jpg']*50))

    async def test_actual_aggregate_rejects_entire_batch(self):
        env=self.env(); env['get_file_temp_object_size'].return_value=15
        paths=[f'1/{"a"*32}/{i}.jpg' for i in range(34)]
        with self.assertRaises(HTTPError):
            await env['image_upload_complete'](types.SimpleNamespace(init_data='', object_paths=paths))
        env['delete_file_temp_batch'].assert_awaited_once_with(paths)
        self.assertFalse(self.spawned)

    async def test_all_complete_endpoints_deduplicate_after_object_deleted(self):
        for kind in ['audio','file','image']:
            env=self.env()
            payload=types.SimpleNamespace(init_data='', object_path=f'1/{"a"*32}.pdf',file_name='x.pdf',object_paths=[f'1/{"a"*32}/0.jpg'])
            if kind=='audio': payload.object_path=f'1/{"a"*32}.mp3';payload.file_name='x.mp3'
            fn=env[kind+'_upload_complete']
            await fn(payload)
            env['get_file_temp_object_size'].return_value=None
            env['get_audio_temp_object_size'].return_value=None
            self.assertTrue((await fn(payload))['duplicate'])
            self.assertEqual(len(self.spawned),1)

    async def test_unsupported_document_rejected_before_target(self):
        env=self.env()
        with self.assertRaises(HTTPError) as exc:
            await env['file_upload_init'](types.SimpleNamespace(init_data='',file_size=99,file_name='x.exe'))
        self.assertEqual(exc.exception.status_code,415)
        env['create_file_upload_target'].assert_not_awaited()


class FormatTests(unittest.TestCase):
    def test_office_archive_regressions(self):
        # Execute the existing preflight unchanged, injecting its production constants.
        env=load_functions('utils.py', {'_validate_office_archive'}, zipfile=zipfile,
            _MAX_OFFICE_ARCHIVE_ENTRIES=10000, _MAX_OFFICE_UNCOMPRESSED_BYTES=512*1024*1024,
            _MAX_OFFICE_SINGLE_ENTRY_BYTES=256*1024*1024,
            _MAX_OFFICE_SUSPICIOUS_RATIO=200, _MAX_OFFICE_RATIO_CHECK_BYTES=50*1024*1024)
        check=env['_validate_office_archive']
        with tempfile.TemporaryDirectory() as directory:
            from docx import Document
            from pptx import Presentation
            doc=Document();doc.add_paragraph('normal document');doc.save(directory+'/x.docx')
            prs=Presentation();prs.slides.add_slide(prs.slide_layouts[0]);prs.save(directory+'/x.pptx')
            check(directory+'/x.docx');check(directory+'/x.pptx')
            Path(directory+'/bad.docx').write_bytes(b'broken ZIP')
            with self.assertRaises(ValueError):check(directory+'/bad.docx')
            with zipfile.ZipFile(directory+'/bomb.docx','w',zipfile.ZIP_DEFLATED) as archive:
                archive.writestr('bomb',b'0'*(51*1024*1024))
            with self.assertRaises(ValueError):check(directory+'/bomb.docx')
            env['_MAX_OFFICE_UNCOMPRESSED_BYTES']=1024
            with self.assertRaises(ValueError):check(directory+'/x.docx')

    def test_pdf_normal_encrypted_malformed(self):
        import fitz
        with tempfile.TemporaryDirectory() as d:
            source=fitz.open()
            for _ in range(150):source.new_page()
            source.save(d+'/ok.pdf');source.save(d+'/encrypted.pdf',encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='password');source.close()
            self.assertEqual(validate_document(d+'/ok.pdf'),150)
            Path(d+'/bad.pdf').write_bytes(b'not a PDF')
            for name in ['encrypted.pdf','bad.pdf']:
                with self.assertRaises(ValueError):validate_document(d+'/'+name)

    def test_image_dimension_guard(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as d:
            Image.new('RGB',(100,100)).save(d+'/ok.png');validate_image(d+'/ok.png')
            Image.new('1',(16001,1)).save(d+'/wide.png')
            with self.assertRaises(ValueError):validate_image(d+'/wide.png')

    def test_bounded_r2_stream_and_partial_cleanup(self):
        env=load_functions('helpers/r2_helper.py', {'_download_object_sync'},os=os,shutil=__import__('shutil'),log_error=Mock(),logger=Mock())
        with tempfile.TemporaryDirectory() as d:
            for declared,body in [(8,b'123456789'),(8,b'1234'),(4,b'1234')]:
                env['_get_r2_client']=lambda:types.SimpleNamespace(get_object=lambda **k:{'ContentLength':declared,'Body':io.BytesIO(body)})
                result=env['_download_object_sync']('bucket','key',d+'/out',8)
                self.assertEqual(result,declared==4)
                self.assertEqual(Path(d+'/out').exists(),result)
