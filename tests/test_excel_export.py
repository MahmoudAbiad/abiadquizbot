"""Check the actual XLSX archive and reopened cell values without app services."""
import ast
import io
import json
import re
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook
from PIL import Image
from reliability_support import load_functions


def exporter():
    path = 'services/export_service.py'
    tree = ast.parse(Path(path).read_text(encoding='utf-8'))
    env = dict(io=io, re=re)
    # Load production constants and pure helpers without image/Telegram imports.
    for node in tree.body:
        if isinstance(node, ast.Assign) and all(isinstance(t, ast.Name) for t in node.targets):
            name = node.targets[0].id
            if name.startswith('STYLE_') or name in {
                'DEFAULT_STYLE', 'VALID_STYLES', '_ARABIC_RE', '_LETTER_RE', '_PARENS_RE',
                '_LATEX_SPAN_RE', 'ARABIC_LETTERS', 'ENGLISH_LETTERS'
            }:
                exec(compile(ast.Module(body=[node], type_ignores=[]), path, 'exec'), env)
    png = io.BytesIO()
    Image.new('RGB', (1500, 1600), 'white').save(png, format='PNG')
    env['iter_print_blocks'] = Mock(side_effect=lambda *args: iter([png.getvalue()]))
    return load_functions(path, {
        'build_quiz_xlsx', '_add_xlsx_math_preview', '_is_math_question',
        '_answer_print_question', '_letters_for',
        'ExportError', 'normalize_style', 'detect_language', '_strip_parens'
    }, **env)


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.env = exporter()
        self.build = self.env['build_quiz_xlsx']
        self.questions = [
            {'question': 'ما ناتج جمع العددين؟', 'options': ['واحد', 'اثنان', 'ثلاثة'],
             'correct_option_id': 1, 'explanation': 'شرح سري للأستاذ فقط'},
            {'question': 'اختر الإجابة التالية', 'options': ['نعم', 'لا'],
             'correct_option_id': '0', 'points': 2},
        ]

    def test_teacher_arabic_and_variable_options(self):
        wb = load_workbook(io.BytesIO(self.build('اختبار', self.questions)))
        self.assertEqual(wb.sheetnames, ['الأسئلة', 'معلومات الاختبار'])
        sheet = wb.worksheets[0]
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:J3')
        self.assertEqual(sheet['H2'].value, 2)  # zero-based source -> one-based Excel
        self.assertEqual(sheet['I2'].value, 'اثنان')
        self.assertEqual(sheet['J2'].value, 'شرح سري للأستاذ فقط')
        self.assertIsNone(sheet['F3'].value)  # padding for a shorter options list
        self.assertEqual(sheet['G3'].value, 2)
        self.assertEqual(wb.worksheets[1]['B2'].value, 'اختبار')

    def test_student_archive_contains_no_answer_key_or_explanation(self):
        data = self.build('اختبار', self.questions, include_answers=False)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.worksheets[0].max_column, 7)
        self.assertEqual(len(wb.sheetnames), 2)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = ''.join(archive.read(n).decode() for n in archive.namelist() if n.endswith('.xml'))
        self.assertNotIn('شرح سري للأستاذ فقط', xml)
        self.assertNotIn('الإجابة الصحيحة', xml)
        self.assertNotIn('state="hidden"', xml)

    def test_formula_and_link_text_remain_literal(self):
        q = {'question': '=HYPERLINK("https://example.com","x")',
             'options': ['+1', '-2', '@SUM(1,2)', 'https://example.com'],
             'correct_option_id': 0, 'explanation': '=1+1'}
        wb = load_workbook(io.BytesIO(self.build('=1+1', [q])))
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'f')
                    self.assertIsNone(cell.hyperlink)
        self.assertEqual(wb.worksheets[0]['B2'].value, q['question'])
        self.assertEqual(wb.worksheets[1]['B2'].value, '=1+1')

    def test_math_tables_and_matrices_preserve_source(self):
        q = self.questions[0] | {
            'question': r'احسب $\frac{1}{2}$',
            'table': {'headers': ['س'], 'rows': [['1']]},
            'matrices': [{'rows': [['1', '2']]}],
        }
        wb = load_workbook(io.BytesIO(self.build('رياضيات', [q])))
        self.assertEqual(wb.active.title, 'العرض الرياضي')
        self.assertEqual(len(wb.active._images), 2)
        sheet = wb['الأسئلة']
        self.assertEqual(sheet['B2'].value, q['question'])
        self.assertEqual(json.loads(sheet['H2'].value), q['table'])
        self.assertEqual(json.loads(sheet['I2'].value), q['matrices'])

    def test_student_math_preview_only_receives_visible_fields(self):
        q = self.questions[0] | {'question': r'احسب $x^2$', 'is_math': True}
        data = self.build('اختبار', [q], include_answers=False)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.active.title, 'العرض الرياضي')
        self.assertEqual(len(wb.active._images), 1)
        renderer = self.env['iter_print_blocks']
        renderer.assert_called_once()
        self.assertEqual(renderer.call_args.args[0], {'question': q['question'], 'options': q['options']})
        self.assertEqual(wb['الأسئلة'].max_column, 7)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertEqual(len([n for n in archive.namelist() if n.startswith('xl/media/')]), 1)
            xml = ''.join(archive.read(n).decode() for n in archive.namelist() if n.endswith('.xml'))
        self.assertNotIn(q['explanation'], xml)

    def test_long_panels_have_disjoint_anchors_and_valid_row_heights(self):
        q = self.questions[0] | {'is_math': True}
        wb = load_workbook(io.BytesIO(self.build('اختبار', [q, q])))
        sheet = wb.active
        self.assertEqual(len(sheet._images), 4)
        for previous, current in zip(sheet._images, sheet._images[1:]):
            start = previous.anchor._from.row
            end = current.anchor._from.row
            reserved_pixels = sum(sheet.row_dimensions[r].height * 96 / 72 for r in range(start + 1, end + 1))
            height_pixels = previous.anchor.ext.cy / 9525
            self.assertGreater(reserved_pixels, height_pixels)
        self.assertTrue(all(d.height <= 409 for d in sheet.row_dimensions.values()))

    def test_math_render_failure_is_explicit(self):
        self.env['iter_print_blocks'].side_effect = ValueError('Invalid math')
        with self.assertRaisesRegex(self.env['ExportError'], 'السؤال 1'):
            self.build('اختبار', [self.questions[0] | {'is_math': True}])

    def test_math_only_in_teacher_explanation_creates_preview(self):
        q = self.questions[0] | {'explanation': r'نستخدم $\\frac{1}{2}$'}
        teacher = load_workbook(io.BytesIO(self.build('اختبار', [q])))
        student = load_workbook(io.BytesIO(self.build('اختبار', [q], include_answers=False)))
        self.assertEqual(teacher.active.title, 'العرض الرياضي')
        self.assertEqual(student.active.title, 'الأسئلة')

    def test_english_and_invalid_answer_index(self):
        for index in (None, -1, 100, 'invalid'):
            q = {'question': 'Select the answer', 'options': ['yes', 'no'], 'correct_option_id': index}
            sheet = load_workbook(io.BytesIO(self.build('Quiz', [q]))).worksheets[0]
            self.assertEqual(sheet.title, 'Questions')
            self.assertFalse(sheet.sheet_view.rightToLeft)
            self.assertIsNone(sheet['G2'].value)
            self.assertIsNone(sheet['H2'].value)

    def test_empty_and_invalid_or_oversized_text_are_explicit_errors(self):
        with self.assertRaises(self.env['ExportError']):
            self.build('Quiz', [])
        for text in ('x' * 32768, 'bad\x00text'):
            with self.assertRaises(self.env['ExportError']):
                self.build('Quiz', [{'question': text, 'options': ['a', 'b']}])


if __name__ == '__main__':
    unittest.main()
