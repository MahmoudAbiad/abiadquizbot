"""اختبارات ميزة ترحيل العمليات الثقيلة (services/offload.py + السيرفر الثاني).

تُشغَّل بدون شبكة ولا FastAPI/httpx: النقل (_post) يُستبدل بدالة تستدعي نواة الـ worker
مباشرة بعد تمرير الـ payload عبر JSON فعلي (لالتقاط أي عدم تطابق بأسماء الحقول بين
الطرفين). طبقة FastAPI تُختبر بـ stubs بسيطة للإطار فقط (منطق _run الحقيقي هو المُختبَر).

التشغيل:  python -m unittest tests.test_offload -v   (من جذر المشروع)
"""
import asyncio
import importlib
import json
import os
import subprocess
import sys
import types
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

MATH_Q = {
    "question": "ما قيمة $x^2 + 2x + 1$ عند $x=3$؟",
    "options": ["16", "9", "12", "10"],
    "correct_option_id": 0,
    "explanation": "تعويض مباشر",
}
TEXT_Q = {"question": "عاصمة فرنسا؟", "options": ["باريس", "روما", "مدريد", "برلين"], "correct_option_id": 0}


def _run_env(env_overrides, code):
    """ينفّذ كوداً بعملية بايثون جديدة ببيئة محددة (الإعدادات تُقرأ عند الاستيراد)."""
    env = {k: v for k, v in os.environ.items() if not k.startswith(("OFFLOAD_", "APP_ROLE"))}
    env.update(env_overrides)
    out = subprocess.run([sys.executable, "-c", code], cwd=ROOT, env=env, capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stderr
    return out.stdout.strip().splitlines()[-1]


class ConfigTests(unittest.TestCase):
    CODE = "from services import offload as o; print(o.ENABLED, o.UPDATE_CONCURRENCY_LIMIT, o.BASE_LOCAL_RENDER_LIMIT)"

    def test_off_by_default_keeps_original_limits(self):
        self.assertEqual(_run_env({}, self.CODE), "False 15 3")

    def test_flag_zero_keeps_original_limits(self):
        env = {"OFFLOAD_ENABLED": "0", "OFFLOAD_WORKER_URL": "https://w.example", "OFFLOAD_WORKER_SECRET": "s"}
        self.assertEqual(_run_env(env, self.CODE), "False 15 3")

    def test_enabled_without_url_or_secret_is_disabled_not_crashing(self):
        self.assertEqual(_run_env({"OFFLOAD_ENABLED": "1"}, self.CODE), "False 15 3")
        self.assertEqual(_run_env({"OFFLOAD_ENABLED": "1", "OFFLOAD_WORKER_URL": "https://w.example"}, self.CODE), "False 15 3")

    def test_enabled_expands_update_limit_only(self):
        env = {"OFFLOAD_ENABLED": "true", "OFFLOAD_WORKER_URL": "https://w.example/", "OFFLOAD_WORKER_SECRET": "s"}
        # سقف الرسم المحلي لا يتوسّع أبداً (ذاكرة الأساسي ثابتة)، سقف التحديثات فقط
        self.assertEqual(_run_env(env, self.CODE), "True 24 3")

    def test_update_limit_never_below_original(self):
        env = {"OFFLOAD_ENABLED": "1", "OFFLOAD_WORKER_URL": "https://w", "OFFLOAD_WORKER_SECRET": "s", "OFFLOAD_UPDATE_CONCURRENCY": "5"}
        self.assertEqual(_run_env(env, self.CODE), "True 15 3")

    def test_worker_role_does_not_import_bot(self):
        # APP_ROLE=render_worker يجب ألا يحمّل config (BOT_TOKEN/Redis): نتحقق ساكناً من ترتيب main.py
        src = (ROOT / "main.py").read_text(encoding="utf-8")
        self.assertLess(src.index('"-m", "render_worker"'), src.index("from config import"))


class BreakerTests(unittest.TestCase):
    def test_old_inflight_result_cannot_close_newly_opened_circuit(self):
        from services.offload import _Breaker
        b = _Breaker(threshold=1)
        old_generation = b.generation
        b.record_failure(old_generation)
        b.record_success(old_generation)
        self.assertIsNotNone(b.opened_at)
        self.assertFalse(b.allow())

    def test_recovery_probe_stays_unique_after_cooldown(self):
        from services.offload import _Breaker
        now = [0.0]
        b = _Breaker(threshold=1, clock=lambda: now[0])
        b.record_failure()
        now[0] = 31
        self.assertTrue(b.allow())
        now[0] = 100
        self.assertFalse(b.allow())
        b.record_success(b.generation)
        self.assertTrue(b.allow())

    def test_opens_after_threshold_then_half_open_after_cooldown(self):
        from services.offload import _Breaker
        now = [0.0]
        b = _Breaker(threshold=3, cooldown=30, clock=lambda: now[0])
        for _ in range(2):
            b.record_failure()
        self.assertTrue(b.allow())
        b.record_failure()
        self.assertFalse(b.allow())
        now[0] = 29
        self.assertFalse(b.allow())
        now[0] = 31
        self.assertTrue(b.allow())      # محاولة واحدة (half-open)
        self.assertFalse(b.allow())     # والبقية تنتظر نتيجتها
        b.record_success()
        self.assertTrue(b.allow())


class ClientTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        from services import offload
        self.o = offload
        self._saved = (offload.ENABLED, offload.FALLBACK_LOCAL, offload._breaker, dict(offload.stats))
        offload.ENABLED = True
        offload.FALLBACK_LOCAL = True
        offload._breaker = offload._Breaker(threshold=3, cooldown=30)
        self.calls = []

    def tearDown(self):
        o = self.o
        o.ENABLED, o.FALLBACK_LOCAL, o._breaker = self._saved[0], self._saved[1], self._saved[2]
        o.stats.clear(); o.stats.update(self._saved[3])

    def _transport(self, *results):
        results = list(results)
        async def fake(path, body, read_timeout):
            self.calls.append((path, json.loads(body.decode("utf-8")), read_timeout))
            r = results.pop(0) if len(results) > 1 else results[0]
            if isinstance(r, Exception):
                raise r
            return r
        return patch.object(self.o, "_post", fake)

    async def test_disabled_never_touches_network(self):
        self.o.ENABLED = False
        with self._transport((200, b"\x89PNG\r\n\x1a\nxx")):
            self.assertIsNone(await self.o.remote_render_image(MATH_Q, 0, 1, True))
        self.assertEqual(self.calls, [])

    async def test_disabled_keeps_local_mode_even_when_fallback_is_disabled(self):
        self.o.ENABLED = False
        self.o.FALLBACK_LOCAL = False
        for _ in range(3):
            self.o._breaker.record_failure()
        with self._transport((500, b"")):
            self.assertIsNone(await self.o.remote_render_image(MATH_Q, 0, 1, True))
            self.assertIsNone(await self.o.remote_build_quiz("pdf", "t", [TEXT_Q], "simple"))
            self.assertIsNone(await self.o.remote_build_document("pdf", "t", "text"))
        self.assertEqual(self.calls, [])

    async def test_queued_requests_recheck_open_circuit(self):
        calls = []
        async def down(*args):
            calls.append(args)
            await asyncio.sleep(0)
            return 503, b""
        with patch.object(self.o, "_remote_sem", asyncio.Semaphore(1)), \
             patch.object(self.o, "_post", down):
            results = await asyncio.gather(*[
                self.o.remote_render_image(MATH_Q, 0, 1, True) for _ in range(20)
            ])
        self.assertEqual(len(calls), 3)
        self.assertTrue(all(r is None for r in results))
        self.assertEqual(self.o.stats["skipped_circuit_open"], 17)
        self.assertEqual(self.o.stats["fallback_local"], 20)

    async def test_open_circuit_honors_disabled_fallback_for_all_operations(self):
        self.o.FALLBACK_LOCAL = False
        for _ in range(3):
            self.o._breaker.record_failure()
        operations = [
            lambda: self.o.remote_render_image(MATH_Q, 0, 1, True),
            lambda: self.o.remote_build_quiz("pdf", "t", [TEXT_Q], "simple"),
            lambda: self.o.remote_build_document("pdf", "t", "text"),
        ]
        with self._transport((500, b"")):
            for operation in operations:
                with self.assertRaises(self.o.OffloadUnavailable):
                    await operation()
        self.assertEqual(self.calls, [])
        self.assertEqual(self.o.stats["fallback_local"], 0)

    async def test_oversized_payload_honors_disabled_fallback(self):
        self.o.FALLBACK_LOCAL = False
        with patch.object(self.o, "MAX_REQUEST_BYTES", 10), self._transport((200, b"%PDF-")):
            with self.assertRaises(self.o.OffloadUnavailable):
                await self.o.remote_build_document("pdf", "t", "text")
        self.assertEqual(self.calls, [])

    async def test_unserializable_payload_honors_fallback_policy(self):
        question = {}
        question["circular"] = question
        with self._transport((200, b"\x89PNG")):
            self.assertIsNone(await self.o.remote_render_image(question, 0, 1, True))
            self.o.FALLBACK_LOCAL = False
            with self.assertRaises(self.o.OffloadUnavailable):
                await self.o.remote_render_image(question, 0, 1, True)
        self.assertEqual(self.calls, [])

    async def test_probe_is_reserved_after_queue_and_cancellation_releases_it(self):
        now = [0.0]
        self.o._breaker = self.o._Breaker(threshold=1, clock=lambda: now[0])
        self.o._breaker.record_failure()
        now[0] = 31
        started = asyncio.Event()
        async def pending(*args):
            started.set()
            await asyncio.Event().wait()
        sem = asyncio.Semaphore(1)
        await sem.acquire()
        with patch.object(self.o, "_remote_sem", sem), patch.object(self.o, "_post", pending):
            request = asyncio.create_task(self.o.remote_render_image(MATH_Q, 0, 1, True))
            await asyncio.sleep(0)
            self.assertFalse(self.o._breaker.probe_in_flight)
            sem.release()
            await asyncio.wait_for(started.wait(), 2)
            self.assertTrue(self.o._breaker.probe_in_flight)
            now[0] = 100
            self.assertIsNone(await self.o.remote_render_image(MATH_Q, 0, 1, True))
            request.cancel()
            with self.assertRaises(asyncio.CancelledError):
                await request
        self.assertFalse(self.o._breaker.probe_in_flight)
        self.assertFalse(self.o.is_active())
        now[0] = 131
        with self._transport((200, b"\x89PNG\r\n\x1a\nrecovered")):
            self.assertIsNotNone(await self.o.remote_render_image(MATH_Q, 0, 1, True))
        self.assertIsNone(self.o._breaker.opened_at)

    async def test_success_returns_bytes_and_payload_shape(self):
        png = b"\x89PNG\r\n\x1a\n" + b"data"
        with self._transport((200, png)):
            self.assertEqual(await self.o.remote_render_image(MATH_Q, 2, 5, True), png)
        path, body, timeout = self.calls[0]
        self.assertEqual(path, "/v1/render/question-image")
        self.assertEqual(body, {"question": MATH_Q, "idx": 2, "total": 5, "is_ar": True})

    async def test_http_error_falls_back_to_local_and_counts(self):
        with self._transport((500, b"boom")):
            self.assertIsNone(await self.o.remote_render_image(MATH_Q, 0, 1, True))
        self.assertEqual(self.o.stats["fallback_local"], 1)

    async def test_garbage_200_body_is_treated_as_failure(self):
        with self._transport((200, b"<html>not a png</html>")):
            self.assertIsNone(await self.o.remote_render_image(MATH_Q, 0, 1, True))

    async def test_circuit_opens_and_stops_calling_worker(self):
        with self._transport((503, b"")):
            for _ in range(3):
                await self.o.remote_render_image(MATH_Q, 0, 1, True)
            n = len(self.calls)
            self.assertIsNone(await self.o.remote_render_image(MATH_Q, 0, 1, True))
        self.assertEqual(len(self.calls), n)   # لم يُرسل طلب رابع
        self.assertEqual(self.o.stats["skipped_circuit_open"], 1)

    async def test_fallback_disabled_raises(self):
        self.o.FALLBACK_LOCAL = False
        with self._transport((500, b"")):
            with self.assertRaises(self.o.OffloadUnavailable):
                await self.o.remote_build_document("pdf", "t", "نص")

    async def test_connect_error_retried_once_then_succeeds(self):
        pdf = b"%PDF-1.4 ..."
        real_sleep = asyncio.sleep
        async def fast_sleep(_): await real_sleep(0)
        with patch.object(self.o.asyncio, "sleep", fast_sleep), \
             self._transport(self.o._TransportError("down", retryable=True), (200, pdf)):
            self.assertEqual(await self.o.remote_build_document("pdf", "t", "نص"), pdf)
        self.assertEqual(len(self.calls), 2)

    async def test_read_timeout_not_retried(self):
        with self._transport(self.o._TransportError("read timeout", retryable=False)):
            self.assertIsNone(await self.o.remote_build_quiz("pdf", "t", [TEXT_Q], "simple"))
        self.assertEqual(len(self.calls), 1)

    async def test_422_becomes_export_error_without_fallback_or_breaker_hit(self):
        from services.export_service import ExportError
        msg = json.dumps({"message": "لا توجد أسئلة لتصديرها."}, ensure_ascii=False).encode()
        with self._transport((422, msg)):
            with self.assertRaises(ExportError) as ctx:
                await self.o.remote_build_quiz("docx", "t", [], "simple")
        self.assertIn("لا توجد أسئلة", str(ctx.exception))
        self.assertEqual(self.o._breaker.failures, 0)

    async def test_oversized_payload_processed_locally(self):
        huge = "ا" * (self.o.MAX_REQUEST_BYTES + 10)
        with self._transport((200, b"%PDF-")):
            self.assertIsNone(await self.o.remote_build_document("pdf", "t", huge))
        self.assertEqual(self.calls, [])

    async def test_invalid_format_never_sent(self):
        with self._transport((200, b"%PDF-")):
            self.assertIsNone(await self.o.remote_build_quiz("exe", "t", [TEXT_Q], "simple"))
        self.assertEqual(self.calls, [])


class WorkerCoreTests(unittest.TestCase):
    def test_student_export_does_not_include_teacher_answers(self):
        import io
        from docx import Document
        from services import offload_worker_core as core
        question = dict(TEXT_Q, explanation='UNIQUE_TEACHER_ANSWER_SENTINEL')
        for include_answers in (False, True):
            data, _ = core.build_quiz({'fmt': 'docx', 'questions': [question],
                                       'include_answers': include_answers})
            document = Document(io.BytesIO(data))
            text = '\n'.join(p.text for p in document.paragraphs)
            text += '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            # Current print exports draw answers as images, not selectable text.
            self.assertEqual(('Answers and solutions' in text or 'الإجابات والحلول' in text), include_answers)
            if include_answers:
                teacher_shapes = len(document.inline_shapes)
            else:
                student_shapes = len(document.inline_shapes)
        self.assertGreater(teacher_shapes, student_shapes)

    def test_include_answers_requires_boolean(self):
        from services import offload_worker_core as core
        with self.assertRaises(core.BadPayload):
            core.build_quiz({'fmt': 'docx', 'questions': [TEXT_Q], 'include_answers': 'false'})

    """يشغّل الدوال الحقيقية (matplotlib/docx/reportlab) كما سيفعل السيرفر الثاني."""

    def test_render_image_png(self):
        from services import offload_worker_core as core
        data, mt = core.render_image({"question": MATH_Q, "idx": 0, "total": 3, "is_ar": True})
        self.assertTrue(data.startswith(b"\x89PNG"))
        self.assertEqual(mt, "image/png")

    def test_build_quiz_docx_and_pdf(self):
        from services import offload_worker_core as core
        docx, mt = core.build_quiz({"fmt": "docx", "title": "اختبار", "questions": [TEXT_Q, MATH_Q], "style": "modern"})
        self.assertTrue(docx.startswith(b"PK"))
        self.assertIn("wordprocessingml", mt)
        pdf, mt = core.build_quiz({"fmt": "pdf", "title": "اختبار", "questions": [TEXT_Q, MATH_Q], "style": "academic"})
        self.assertTrue(pdf.startswith(b"%PDF-"))

    def test_build_document(self):
        from services import offload_worker_core as core
        text = "# عنوان\n\n- نقطة أولى\n- نقطة ثانية\n\nفقرة عادية."
        for fmt, magic in (("docx", b"PK"), ("pdf", b"%PDF-")):
            data, _ = core.build_document({"fmt": fmt, "title": "محاضرة", "text": text})
            self.assertTrue(data.startswith(magic))

    def test_unknown_style_falls_back_to_default(self):
        from services import offload_worker_core as core
        data, _ = core.build_quiz({"fmt": "pdf", "title": "t", "questions": [TEXT_Q], "style": "hacker"})
        self.assertTrue(data.startswith(b"%PDF-"))

    def test_export_error_propagates(self):
        from services import offload_worker_core as core
        from services.export_service import ExportError
        with self.assertRaises(ExportError):
            core.build_quiz({"fmt": "pdf", "title": "t", "questions": []})
        with self.assertRaises(ExportError):
            core.build_document({"fmt": "docx", "title": "t", "text": "   "})

    def test_payload_validation(self):
        from services import offload_worker_core as core
        bad = [
            (core.render_image, {"question": "x", "idx": 0, "total": 1, "is_ar": True}),
            (core.render_image, {"question": MATH_Q, "idx": -1, "total": 1, "is_ar": True}),
            (core.render_image, {"question": MATH_Q, "idx": 0, "total": 1, "is_ar": "yes"}),
            (core.render_image, {"question": MATH_Q, "idx": True, "total": 1, "is_ar": True}),
            (core.build_quiz, {"fmt": "exe", "title": "t", "questions": [TEXT_Q]}),
            (core.build_quiz, {"fmt": "pdf", "title": "t", "questions": ["not a dict"]}),
            (core.build_quiz, {"fmt": "pdf", "title": "t", "questions": [TEXT_Q] * (core.MAX_QUESTIONS + 1)}),
            (core.build_quiz, {"fmt": "pdf", "title": "x" * (core.MAX_TITLE_CHARS + 1), "questions": [TEXT_Q]}),
            (core.build_document, {"fmt": "pdf", "title": "t", "text": "x" * (core.MAX_TEXT_CHARS + 1)}),
        ]
        for fn, payload in bad:
            with self.assertRaises(core.BadPayload, msg=f"{fn.__name__} {str(payload)[:60]}"):
                fn(payload)


class EndToEndNoNetworkTests(unittest.IsolatedAsyncioTestCase):
    """العميل -> JSON حقيقي -> نواة الـ worker الحقيقية -> bytes -> العميل. بدون HTTP فقط."""

    def setUp(self):
        from services import offload
        self.o = offload
        self._saved = (offload.ENABLED, offload._breaker, dict(offload.stats))
        offload.ENABLED = True
        offload._breaker = offload._Breaker()
        for key in offload.stats:
            offload.stats[key] = 0

        from services import offload_worker_core as core
        from services.export_service import ExportError
        routes = {"/v1/render/question-image": core.render_image, "/v1/export/quiz": core.build_quiz,
                  "/v1/export/document": core.build_document}

        async def fake_post(path, body, read_timeout):
            payload = json.loads(body.decode("utf-8"))     # نفس ما يفعله السيرفر الثاني
            try:
                data, _ = await asyncio.to_thread(routes[path], payload)
                return 200, data
            except ExportError as exc:
                return 422, json.dumps({"message": str(exc)}, ensure_ascii=False).encode()
        self.patch = patch.object(offload, "_post", fake_post)
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.o.ENABLED, self.o._breaker = self._saved[0], self._saved[1]
        self.o.stats.clear(); self.o.stats.update(self._saved[2])

    async def test_all_three_endpoints_round_trip(self):
        png = await self.o.remote_render_image(MATH_Q, 0, 2, True)
        self.assertTrue(png.startswith(b"\x89PNG"))
        docx = await self.o.remote_build_quiz("docx", "اختبار", [TEXT_Q, MATH_Q], "simple")
        self.assertTrue(docx.startswith(b"PK"))
        pdf = await self.o.remote_build_document("pdf", "محاضرة", "نص تجريبي طويل قليلاً.")
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertEqual(self.o.stats["remote_ok"] >= 3, True)

    async def test_render_async_uses_remote_then_local_when_off(self):
        from services import image_quiz_renderer as r
        calls = []
        real = self.o.remote_render_image
        async def spy(*a):
            calls.append(a)
            return await real(*a)
        with patch.object(self.o, "remote_render_image", spy):
            out = await r.render_question_image_async(MATH_Q, 0, 1, True)
            self.assertTrue(out.startswith(b"\x89PNG"))
            self.assertEqual(len(calls), 1)
            self.assertEqual(self.o.stats["fallback_local"], 0)

        # الميزة مغلقة => المسار المحلي القديم (لا استدعاء شبكي)
        self.o.ENABLED = False
        out = await r.render_question_image_async(MATH_Q, 0, 1, True)
        self.assertTrue(out.startswith(b"\x89PNG"))

    async def test_remote_down_falls_back_to_local_render(self):
        from services import image_quiz_renderer as r
        async def down(path, body, read_timeout):
            raise self.o._TransportError("worker asleep", retryable=False)
        with patch.object(self.o, "_post", down):
            out = await r.render_question_image_async(MATH_Q, 0, 1, True)
        self.assertTrue(out.startswith(b"\x89PNG"))      # المستخدم لم يتأثر
        self.assertGreaterEqual(self.o.stats["fallback_local"], 1)


# ---------- طبقة FastAPI (render_worker.py) بإطار مُستبدل بـ stubs ----------

def _install_fastapi_stubs():
    fastapi = types.ModuleType("fastapi")

    class HTTPException(Exception):
        def __init__(self, status_code, detail=None, headers=None):
            self.status_code, self.detail, self.headers = status_code, detail, headers

    class FastAPI:
        def __init__(self, *a, **k): pass
        def api_route(self, *a, **k): return lambda f: f
        def post(self, *a, **k): return lambda f: f

    fastapi.FastAPI, fastapi.HTTPException, fastapi.Request = FastAPI, HTTPException, object
    responses = types.ModuleType("fastapi.responses")

    class Response:
        def __init__(self, content=b"", media_type=None, status_code=200):
            self.body, self.media_type, self.status_code = content, media_type, status_code

    class JSONResponse(Response):
        def __init__(self, status_code=200, content=None):
            super().__init__(json.dumps(content, ensure_ascii=False).encode(), "application/json", status_code)

    responses.Response, responses.JSONResponse = Response, JSONResponse
    sys.modules.setdefault("fastapi", fastapi)
    sys.modules.setdefault("fastapi.responses", responses)


class FakeRequest:
    def __init__(self, payload, secret="s3cret", headers=None, raw=None):
        self.headers = {"x-worker-secret": secret} if secret is not None else {}
        self.headers.update(headers or {})
        self._raw = raw if raw is not None else json.dumps(payload, ensure_ascii=False).encode()

    async def body(self):
        return self._raw


class WorkerHttpLayerTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        _install_fastapi_stubs()
        os.environ["OFFLOAD_WORKER_SECRET"] = "s3cret"
        sys.modules.pop("render_worker", None)
        self.w = importlib.import_module("render_worker")
        self.w.WORKER_SECRET = "s3cret"
        self.w._waiting = 0
        # HTTP admission tests use a thread substitute; real spawned processes
        # and renderer overlap are verified separately in test_render_processes.
        async def submit(work, payload):
            return await asyncio.to_thread(self.w.execute, work, payload)
        self.w._submit = submit

    async def test_auth_required_and_constant_time_compare(self):
        for req in (FakeRequest({}, secret=None), FakeRequest({}, secret="wrong")):
            with self.assertRaises(sys.modules["fastapi"].HTTPException) as ctx:
                await self.w.render_question_image_endpoint(req)
            self.assertEqual(ctx.exception.status_code, 403)

    async def test_fail_closed_when_secret_unset(self):
        self.w.WORKER_SECRET = ""
        with self.assertRaises(sys.modules["fastapi"].HTTPException) as ctx:
            await self.w.export_quiz_endpoint(FakeRequest({}, secret=""))
        self.assertEqual(ctx.exception.status_code, 503)

    async def test_render_ok_returns_png(self):
        resp = await self.w.render_question_image_endpoint(
            FakeRequest({"question": MATH_Q, "idx": 0, "total": 1, "is_ar": True}))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.media_type, "image/png")
        self.assertTrue(resp.body.startswith(b"\x89PNG"))

    async def test_export_error_maps_to_422_and_releases_slot(self):
        before = self.w._slots._value
        resp = await self.w.export_quiz_endpoint(FakeRequest({"fmt": "pdf", "title": "t", "questions": []}))
        self.assertEqual(resp.status_code, 422)
        self.assertIn("لا توجد أسئلة", json.loads(resp.body)["message"])
        self.assertEqual(self.w._slots._value, before)

    async def test_bad_payload_400_and_slot_released(self):
        before = self.w._slots._value
        with self.assertRaises(sys.modules["fastapi"].HTTPException) as ctx:
            await self.w.export_quiz_endpoint(FakeRequest({"fmt": "exe", "title": "t", "questions": []}))
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertEqual(self.w._slots._value, before)

    async def test_invalid_json_and_oversize(self):
        H = sys.modules["fastapi"].HTTPException
        with self.assertRaises(H) as ctx:
            await self.w.export_quiz_endpoint(FakeRequest(None, raw=b"{not json"))
        self.assertEqual(ctx.exception.status_code, 400)
        with self.assertRaises(H) as ctx:
            await self.w.export_quiz_endpoint(FakeRequest({}, headers={"content-length": str(self.w.MAX_BODY_BYTES + 1)}))
        self.assertEqual(ctx.exception.status_code, 413)

    async def test_busy_when_queue_full_or_wait_times_out(self):
        H = sys.modules["fastapi"].HTTPException
        self.w._waiting = self.w.MAX_QUEUE
        self.w._admitted = self.w.CONCURRENCY + self.w.MAX_QUEUE
        with self.assertRaises(H) as ctx:
            await self.w.export_quiz_endpoint(FakeRequest({"fmt": "pdf", "title": "t", "questions": [TEXT_Q]}))
        self.assertEqual(ctx.exception.status_code, 503)
        self.w._admitted = 0

        self.w._waiting = 0
        self.w.QUEUE_WAIT_SECONDS = 0.05
        held = [await self.w._slots.acquire() for _ in range(self.w.CONCURRENCY)]   # كل الخانات مشغولة
        try:
            with self.assertRaises(H) as ctx:
                await self.w.export_quiz_endpoint(FakeRequest({"fmt": "pdf", "title": "t", "questions": [TEXT_Q]}))
            self.assertEqual(ctx.exception.status_code, 503)
            self.assertEqual(self.w._waiting, 0)      # العدّاد لا يتسرّب
        finally:
            for _ in held:
                self.w._slots.release()

    async def test_concurrency_never_exceeds_limit(self):
        active, peak = [0], [0]
        def slow(payload):
            import time
            active[0] += 1; peak[0] = max(peak[0], active[0])
            time.sleep(0.05)
            active[0] -= 1
            return b"%PDF-x", "application/pdf"
        reqs = [self.w._run(FakeRequest({"fmt": "pdf"}), slow) for _ in range(self.w.CONCURRENCY * 3)]
        await asyncio.gather(*reqs)
        self.assertLessEqual(peak[0], self.w.CONCURRENCY)
        self.assertEqual(self.w._slots._value, self.w.CONCURRENCY)


if __name__ == "__main__":
    unittest.main()
