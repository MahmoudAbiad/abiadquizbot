"""Real workbook round trips plus direct Telegram intake without AI or wallet use."""
import asyncio
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch

from openpyxl import Workbook, load_workbook
from services.excel_import import ExcelImportError, read_quiz_xlsx
from reliability_support import load_functions
from test_excel_export import exporter


class ExcelImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'quiz.xlsx'
        self.headers = ['نص السؤال', 'الخيار 1', 'الخيار 2', 'رقم الإجابة الصحيحة']

    def write(self, rows, headers=None):
        book = Workbook()
        book.active.append(headers or self.headers)
        for row in rows:
            book.active.append(row)
        book.save(self.path)
        book.close()

    def test_arabic_numeric_options_and_blank_rows(self):
        self.write([['كم؟', 0, 2, '٢'], [], ['لماذا؟', 'نعم', 'لا', 1]])
        title, questions = read_quiz_xlsx(self.path, 'عنوان')
        self.assertEqual(title, 'عنوان')
        self.assertEqual(len(questions), 2)
        self.assertEqual(questions[0]['options'], ['0', '2'])
        self.assertEqual(questions[0]['correct_option_id'], 1)

    def test_english_text_answer_with_reordered_columns(self):
        self.write([['B', 'B', 'Pick', 'A']], ['Option 2', 'Correct Answer', 'Question', 'Option 1'])
        self.assertEqual(read_quiz_xlsx(self.path)[1][0]['correct_option_id'], 1)

    def test_real_teacher_export_roundtrip_and_edits(self):
        source = [{'question': 'ما الناتج؟', 'options': ['=1+1', '2', 'ثلاثة'],
                   'correct_option_id': 1, 'explanation': 'شرح', 'points': 3}]
        self.path.write_bytes(exporter()['build_quiz_xlsx']('عنوان الاختبار', source))
        book = load_workbook(self.path)
        book['الأسئلة']['B2'] = 'سؤال معدّل'
        book.save(self.path)
        book.close()
        title, questions = read_quiz_xlsx(self.path)
        self.assertEqual(title, 'عنوان الاختبار')
        self.assertEqual(questions[0]['question'], 'سؤال معدّل')
        for key in ('options', 'correct_option_id', 'explanation', 'points'):
            self.assertEqual(questions[0][key], source[0][key])

    def test_math_export_preview_first_preserves_editable_source(self):
        source = [{'question': r'Find $\frac{1}{2}$', 'options': ['$0.5$', '$2$'],
                   'correct_option_id': 0, 'is_math': True,
                   'table': {'headers': ['x'], 'rows': [['1']]},
                   'matrices': [{'rows': [['1', '0'], ['0', '1']], 'bracket': 'square'}]}]
        self.path.write_bytes(exporter()['build_quiz_xlsx']('Math', source))
        title, questions = read_quiz_xlsx(self.path)
        self.assertEqual(title, 'Math')
        for key in ('question', 'options', 'table', 'matrices', 'is_math', 'correct_option_id'):
            self.assertEqual(questions[0][key], source[0][key])

    def test_invalid_rows_are_actionable_and_atomic(self):
        for row in (['Q', 'a', 'b', 0], ['Q', 'a', 'b', 3], ['Q', 'a', 'b', 1.5],
                    ['Q', 'a', 'b', None], ['', 'a', 'b', 1], ['Q', '', 'b', 1],
                    ['Q', '=1+1', 'b', 1], ['Q', 'a', 'b', '=1+1']):
            with self.subTest(row=row):
                self.write([['valid', 'a', 'b', 1], row])
                with self.assertRaisesRegex(ExcelImportError, 'الصف 3'):
                    read_quiz_xlsx(self.path)

    def test_stale_answer_text_rejected(self):
        self.write([['Q', 'a', 'b', 1, 'b']], self.headers + ['الإجابة الصحيحة'])
        with self.assertRaisesRegex(ExcelImportError, 'لا يطابق'):
            read_quiz_xlsx(self.path)

    def test_student_export_missing_answer_key(self):
        self.path.write_bytes(exporter()['build_quiz_xlsx']('Q', [
            {'question': 'Q', 'options': ['A', 'B'], 'correct_option_id': 0}], include_answers=False))
        with self.assertRaisesRegex(ExcelImportError, 'نسخة الأستاذ'):
            read_quiz_xlsx(self.path)

    def test_duplicate_headers_and_option_gaps(self):
        for headers in (self.headers + ['السؤال'], ['Question', 'Option 1', 'Option 3', 'Correct Answer']):
            self.write([['Q', 'a', 'b', 1]], headers)
            with self.assertRaises(ExcelImportError):
                read_quiz_xlsx(self.path)

    def test_invalid_structured_math_rejected(self):
        for extra in ('null', '[1]', '{"rows":[[1],[1,2]]}', '{"rows":[[{}]]}'):
            self.write([['Q', 'a', 'b', 1, extra]], self.headers + ['الجدول (JSON)'])
            with self.assertRaisesRegex(ExcelImportError, 'JSON'):
                read_quiz_xlsx(self.path)

    def test_limits_corruption_and_empty_sheet(self):
        self.write([['Q', 'a', 'b', 1]] * 121)
        with self.assertRaisesRegex(ExcelImportError, '120'):
            read_quiz_xlsx(self.path)
        self.write([])
        with self.assertRaisesRegex(ExcelImportError, 'فارغة'):
            read_quiz_xlsx(self.path)
        self.path.write_bytes(b'not an Excel archive')
        with self.assertRaisesRegex(ExcelImportError, '.xlsx'):
            read_quiz_xlsx(self.path)

    def test_archive_size_limit(self):
        self.write([['Q', 'a', 'b', 1]])
        with patch('services.excel_import.MAX_ARCHIVE_BYTES', 10):
            with self.assertRaisesRegex(ExcelImportError, 'كبير'):
                read_quiz_xlsx(self.path)


class ExcelIntakeTests(unittest.IsolatedAsyncioTestCase):
    async def test_import_resets_stale_charge_state_and_starts_without_ai(self):
        question = {'question': 'Q', 'options': ['a', 'b'], 'correct_option_id': 0}
        reader = Mock(return_value=('Title', [question]))
        begin = AsyncMock()
        env = load_functions('handlers/files.py', {'_import_excel_quiz'},
                             asyncio=asyncio, read_quiz_xlsx=reader, _start_loaded_quiz=begin)
        state = SimpleNamespace(set_data=AsyncMock())
        message = SimpleNamespace(answer=AsyncMock())
        await env['_import_excel_quiz'](message, state, 'path.xlsx', 'Filename')
        state.set_data.assert_awaited_once_with({'input_type': 'excel', 'content_source': 'xlsx', 'input_channel': 'telegram'})
        begin.assert_awaited_once_with(message, state, [question], 'Title', origin='excel_import')

    async def test_upload_routes_xlsx_before_generic_generation_and_cleans_file(self):
        importer, finalize, clean = AsyncMock(), AsyncMock(), Mock()
        env = load_functions('handlers/files.py', {'handle_media'}, asyncio=asyncio,
            offer_transition=AsyncMock(return_value=False), message_action=Mock(),
            QuizState=SimpleNamespace(processing_file_quiz='processing'), PENDING_REQUEST_STATES=(),
            ensure_directory_exists=Mock(), DOWNLOADS_DIR='downloads', os=os,
            TELEGRAM_IMAGE_DOCUMENT_EXTENSIONS=set(), DOCUMENT_EXTENSIONS={'.pdf'},
            validate_file_size=Mock(return_value=(True, None)), supported_extension=Mock(return_value='.xlsx'),
            uuid=SimpleNamespace(uuid4=lambda: SimpleNamespace(hex='unique')), bot=SimpleNamespace(download=AsyncMock()),
            _import_excel_quiz=importer, _finalize_media_processing=finalize,
            safe_file_cleanup=clean, ExcelImportError=ExcelImportError, log_error=Mock(), logger=Mock())
        message = SimpleNamespace(document=SimpleNamespace(file_name='quiz.xlsx', file_size=10),
                                  from_user=SimpleNamespace(id=42), chat=SimpleNamespace(id=42),
                                  photo=None, answer=AsyncMock())
        state = SimpleNamespace(get_state=AsyncMock(return_value=None))
        await env['handle_media'](message, state)
        importer.assert_awaited_once()
        finalize.assert_not_awaited()
        clean.assert_called_once_with(os.path.join('downloads', '42_unique.xlsx'))
        importer.side_effect = ExcelImportError('الصف 3: أضف الإجابة')
        await env['handle_media'](message, state)
        message.answer.assert_awaited_with('❌ الصف 3: أضف الإجابة', parse_mode=None)


if __name__ == '__main__':
    unittest.main()
