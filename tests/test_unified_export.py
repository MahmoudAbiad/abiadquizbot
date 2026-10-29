import io
import unittest

import fitz
from docx import Document
from docx.oxml.ns import qn

from services.export_service import build_quiz_docx, build_quiz_pdf
from services.quiz_print_layout import option_columns, text_parts


def question(options=None):
    return {"question": "Complete the sentence: I have worked here ____ five years.\n"
                        "أكمل الجملة: أعمل هنا منذ خمس سنوات.",
            "options": options or ["for (لمدة)", "since (منذ)", "during (خلال)", "from (من)"],
            "correct_option_id": 0, "explanation": "Use for with a duration."}


class UnifiedExportTests(unittest.TestCase):
    def test_parts_preserve_lines_and_legacy_translation(self):
        self.assertEqual(text_parts("English question\nسؤال عربي", True),
                         [("English question", False), ("سؤال عربي", True)])
        self.assertEqual(text_parts("English (ترجمة عربية)", True, True),
                         [("English", False), ("ترجمة عربية", True)])
        self.assertEqual(text_parts("English (x + y)", True, True), [("English (x + y)", False)])
        self.assertEqual(text_parts("English (ترجمة with English)", True, True),
                         [("English (ترجمة with English)", False)])

    def test_column_selection_is_adaptive_and_handles_variable_counts(self):
        measure = lambda text, rtl: len(text)
        for options, expected in ((["a"] * 4, 4), (["a" * 10] * 4, 2),
                                  (["a" * 40] * 4, 1), (["a"] * 3, 2),
                                  (["a"] * 5, 4), (["a\nb"] * 4, 1), ([], 1)):
            self.assertEqual(option_columns([text_parts(o) for o in options], 100, measure), expected)

    def test_word_native_grid_and_directions_in_all_styles(self):
        for style in ("simple", "modern", "academic"):
            for answers in (False, True):
                with self.subTest(style=style, answers=answers):
                    doc = Document(io.BytesIO(build_quiz_docx("Test", [question()], style, answers)))
                    self.assertEqual(len(doc.inline_shapes), 0)
                    self.assertEqual(len(doc.tables), 1)
                    self.assertEqual(len(doc.tables[0].columns), 4)
                    self.assertIn("for (لمدة)", doc.tables[0].cell(0, 0).text)
                    en = next(p for p in doc.paragraphs if "Complete the sentence" in p.text)
                    ar = next(p for p in doc.paragraphs if "أكمل الجملة" in p.text)
                    self.assertEqual(en._p.pPr.find(qn("w:bidi")).get(qn("w:val")), "0")
                    self.assertEqual(ar._p.pPr.find(qn("w:bidi")).get(qn("w:val")), "1")
                    self.assertTrue(en.paragraph_format.keep_with_next)
                    self.assertTrue(ar.paragraph_format.keep_with_next)
                    texts = "\n".join(p.text for p in doc.paragraphs)
                    self.assertEqual("Use for with a duration." in texts, answers)

    def test_arabic_grid_orders_choices_from_the_right(self):
        q = question(["واحد", "اثنان", "ثلاثة", "أربعة"])
        q["question"] = "اختر الإجابة الصحيحة"
        doc = Document(io.BytesIO(build_quiz_docx("اختبار", [q], include_answers=False)))
        table = doc.tables[0]
        self.assertIsNotNone(table._tbl.tblPr.find(qn("w:bidiVisual")))
        self.assertEqual(table.cell(0, 0).text, "أ) واحد")

    def test_answer_labels_match_each_questions_language(self):
        english = question()
        arabic = {"question": "اختر الإجابة الصحيحة من الخيارات التالية التي توضح معنى الجملة",
                  "options": ["الإجابة الأولى", "الإجابة الثانية", "الإجابة الثالثة", "الإجابة الرابعة"],
                  "correct_option_id": 1}
        doc = Document(io.BytesIO(build_quiz_docx("اختبار", [english, arabic])))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Correct answer: A)", text)
        self.assertIn("الإجابة الصحيحة: ب)", text)
        with fitz.open(stream=build_quiz_pdf("اختبار", [english, arabic]), filetype="pdf") as pdf:
            self.assertIn("Correct answer: A)", pdf[-1].get_text())

    def test_medium_long_and_more_than_four_choices_in_word(self):
        cases = [("a medium grammar choice", 2),
                 ("A long explanation of the grammar rule that needs the full printable width " * 3, 1)]
        for option, columns in cases:
            doc = Document(io.BytesIO(build_quiz_docx("Test", [question([option] * 5)], include_answers=False)))
            self.assertEqual(len(doc.tables[0].columns), columns)
            actual = [cell.text for row in doc.tables[0].rows for cell in row.cells if cell.text]
            self.assertEqual(len(actual), 5)

    def test_pdf_text_grid_and_question_grouping(self):
        for style in ("simple", "modern", "academic"):
            qs = [question([f"Choice{n}{i}" for i in range(4)]) for n in range(16)]
            for n, q in enumerate(qs):
                q["question"] = f"Prompt{n} Complete the sentence ____ today.\nأكمل الجملة التالية."
            with fitz.open(stream=build_quiz_pdf("Test", qs, style, False), filetype="pdf") as pdf:
                self.assertGreater(len(pdf), 1)
                texts = [page.get_text() for page in pdf]
                for n in range(len(qs)):
                    locations = [i for i, text in enumerate(texts) if f"Prompt{n} " in text]
                    self.assertEqual(len(locations), 1)
                    for i in range(4):
                        self.assertIn(f"Choice{n}{i}", texts[locations[0]])
                for page in pdf:
                    self.assertEqual(len(page.get_images()), 0)
                    for word in page.get_text("words"):
                        self.assertGreaterEqual(word[0], 45)
                        self.assertLessEqual(word[2], page.rect.width - 45)
                        self.assertGreaterEqual(word[1], 20)
                        self.assertLessEqual(word[3], page.rect.height - 15)

    def test_oversized_question_and_option_paginate_without_overflow(self):
        q = question(["unbroken" * 1600, "last choice"])
        q["question"] = "Long question " * 500
        with fitz.open(stream=build_quiz_pdf("Test", [q], include_answers=False), filetype="pdf") as pdf:
            self.assertGreater(len(pdf), 2)
            self.assertIn("last choice", pdf[-1].get_text())
            for page in pdf:
                for word in page.get_text("words"):
                    self.assertGreaterEqual(word[0], 45)
                    self.assertLessEqual(word[2], page.rect.width - 45)
                    self.assertGreaterEqual(word[1], 20)
                    self.assertLessEqual(word[3], page.rect.height - 15)


if __name__ == "__main__":
    unittest.main()
